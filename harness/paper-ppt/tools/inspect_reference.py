"""Extract inspectable design evidence from a reference PPTX.

This is evidence for a human/agent design decision, not an automatic template.
"""
import argparse
import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from zipfile import ZipFile
from lxml import etree


def shapes_with_bounds(shapes, transform=(1, 1, 0, 0)):
    """Map group child coordinates to slide coordinates; rotations need visual QA."""
    sx, sy, dx, dy = transform
    for shape in shapes:
        if shape.left is None or shape.width is None:   # inherited placeholder geometry
            continue
        bounds = (shape.left * sx + dx, shape.top * sy + dy,
                  shape.width * sx, shape.height * sy)
        yield shape, bounds
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.find(qn('p:grpSpPr')).find(qn('a:xfrm'))
            off, ext = xfrm.find(qn('a:chOff')), xfrm.find(qn('a:chExt'))
            child_sx = shape.width / max(1, int(ext.get('cx')))
            child_sy = shape.height / max(1, int(ext.get('cy')))
            yield from shapes_with_bounds(shape.shapes, (
                sx * child_sx, sy * child_sy,
                dx + sx * (shape.left - int(off.get('x')) * child_sx),
                dy + sy * (shape.top - int(off.get('y')) * child_sy)))


def inspect(path):
    prs = Presentation(path)
    colors, fonts, sizes = Counter(), Counter(), Counter()
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        kinds = Counter()
        boxes = []
        total = 0
        for shape, bounds in shapes_with_bounds(slide.shapes):
            total += 1
            kinds[str(shape.shape_type)] += 1
            box = [round(bounds[0] / prs.slide_width, 3),
                   round(bounds[1] / prs.slide_height, 3),
                   round(bounds[2] / prs.slide_width, 3),
                   round(bounds[3] / prs.slide_height, 3)]
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            if shape.has_chart:
                boxes.append({'box': box, 'name': shape.name, 'chart': str(shape.chart.chart_type)})
            for frame in frames:
                if not frame.text.strip():
                    continue
                boxes.append({'box': box, 'name': shape.name, 'text': frame.text[:200]})
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
                        if run.font.size:
                            sizes[round(run.font.size.pt)] += 1
                        try:
                            if run.font.color and run.font.color.rgb:
                                colors[str(run.font.color.rgb)] += 1
                        except (AttributeError, TypeError, ValueError):
                            pass
            try:
                if shape.fill.type and shape.fill.fore_color.rgb:
                    colors[str(shape.fill.fore_color.rgb)] += 1
            except (AttributeError, TypeError, ValueError):
                pass
        slides.append({'slide': number, 'shape_types': dict(kinds),
                       'text_boxes': boxes, 'shape_count': total})
    layouts = []
    for layout in prs.slide_layouts:
        deco = []
        for owner in (layout.slide_master, layout):
            for shape, b in shapes_with_bounds(owner.shapes):
                if shape.is_placeholder:
                    continue
                deco.append({'from': 'master' if owner is layout.slide_master else 'layout', 'name': shape.name,
                             'box_1600': [round(b[0] / prs.slide_width * 1600), round(b[1] / prs.slide_height * 900),
                                          round(b[2] / prs.slide_width * 1600), round(b[3] / prs.slide_height * 900)]})
        used = sum(1 for s in prs.slides if s.slide_layout == layout)
        layouts.append({'name': layout.name, 'placeholders': len(layout.placeholders),
                        'used_by_slides': used, 'decorations': deco})
    theme_evidence = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith('ppt/theme/theme') and name.endswith('.xml'):
                root = etree.fromstring(archive.read(name))
                theme_evidence[name] = {
                    'colors': root.xpath('//*[local-name()="clrScheme"]//*[local-name()="srgbClr"]/@val'),
                    'fonts': root.xpath('//*[local-name()="fontScheme"]//*[@typeface]/@typeface'),
                }
    return {'reference': str(Path(path).resolve()), 'slide_count': len(prs.slides),
            'aspect_ratio': round(prs.slide_width / prs.slide_height, 4),
            'common_colors': colors.most_common(12),
            'common_fonts': fonts.most_common(8),
            'common_font_sizes_pt': sizes.most_common(12), 'slides': slides,
            'themes': theme_evidence, 'layouts': layouts,
            'limitations': 'Counts are explicit declarations, not a visual ranking; inherited styles and rotated objects require rendered inspection.'}


def main():
    p = argparse.ArgumentParser()
    p.add_argument('reference', type=Path)
    p.add_argument('output', type=Path)
    a = p.parse_args()
    result = inspect(a.reference)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    a.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'inspected {result["slide_count"]} slides: {a.output}')


if __name__ == '__main__':
    main()
