"""Fail closed when a generated deck lacks its reference and slide design record."""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from zipfile import ZipFile
from lxml import etree


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def validate(deck_path, plan_path):
    plan = json.loads(Path(plan_path).read_text(encoding='utf-8'))
    if not isinstance(plan, dict):
        return ['plan must be a JSON object']
    prs = Presentation(deck_path)
    errors = []
    if not plan.get('paper'):
        errors.append('plan.paper is required')
    if not plan.get('reference'):
        errors.append('plan.reference is required')
    slides = plan.get('slides', [])
    if not isinstance(slides, list):
        return ['plan.slides must be a list']
    if len(slides) != len(prs.slides):
        errors.append(f'plan has {len(slides)} slides, deck has {len(prs.slides)}')
    for number, slide in enumerate(prs.slides, 1):
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            errors.append(f'slide {number}: speaker notes missing')
        entry = slides[number - 1] if number <= len(slides) else {}
        if not isinstance(entry, dict):
            errors.append(f'slide {number}: plan entry must be an object')
            entry = {}
        for key in ('purpose', 'source', 'reference_slide', 'adaptation'):
            if not entry.get(key):
                errors.append(f'slide {number}: missing plan.{key}')
        shapes = list(walk_shapes(slide.shapes))
        if not any(s.has_text_frame and s.text.strip() for s in shapes):
            errors.append(f'slide {number}: no editable text')
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                coverage = shape.width * shape.height / (prs.slide_width * prs.slide_height)
                if coverage > 0.8:
                    errors.append(f'slide {number}: full-slide picture ({coverage:.0%})')
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size and run.font.size.pt < 16:
                            errors.append(f'slide {number}: text below 16 pt')
    # Read explicit font declarations independently of Canvas, including table
    # cells, grouped objects, and chart labels. Inherited sizes still need render QA.
    with ZipFile(deck_path) as archive:
        for name in archive.namelist():
            if name.startswith(('ppt/slides/slide', 'ppt/charts/chart')) and name.endswith('.xml'):
                root = etree.fromstring(archive.read(name))
                small = [int(value) / 100 for value in root.xpath('//@sz') if str(value).isdigit() and int(value) < 1600]
                if small:
                    errors.append(f'{name}: explicit font below 16 pt (min {min(small):g})')
    return list(dict.fromkeys(errors))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('deck', type=Path)
    parser.add_argument('plan', type=Path)
    a = parser.parse_args()
    errors = validate(a.deck, a.plan)
    for error in errors:
        print('FAIL:', error)
    if errors:
        raise SystemExit(1)
    print(f'PASS: {len(Presentation(a.deck).slides)} slides have a design record and editable content')


if __name__ == '__main__':
    main()
