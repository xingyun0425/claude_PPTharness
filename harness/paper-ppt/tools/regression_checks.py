"""Check meaningful failure cases without automating a visual judgment."""
import json
import tempfile
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt
from PIL import Image
from validate_deck import validate
from render_review import audit_layout, check_review, sha256
from inspect_reference import inspect
from deckkit import Canvas, luminance, shift_hue, shift_svg_hue


def main():
    results = []
    def record(name, condition):
        if not condition:
            raise AssertionError(name)
        results.append(name)
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        deck, plan = folder / "deck.pptx", folder / "plan.json"
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        group = s.shapes.add_group_shape()
        box = group.shapes.add_textbox(100000, 100000, 3000000, 800000)
        r = box.text_frame.paragraphs[0].add_run()
        r.text, r.font.size = "Grouped native text", Pt(20)
        s.notes_slide.notes_text_frame.text = "Synthetic fixture source"
        prs.save(deck)
        entry = {"purpose": "test", "source": "fixture", "reference_slide": "1", "adaptation": "grouped text"}
        plan.write_text(json.dumps({"paper": "synthetic fixture", "reference": "fixture.pptx", "slides": [entry]}), encoding="utf-8")
        record("plan without an illustration field is rejected", any("plan.illustration" in x for x in validate(deck, plan)))
        entry["illustration"] = "none"
        plan.write_text(json.dumps({"paper": "synthetic fixture", "reference": "fixture.pptx", "slides": [entry]}), encoding="utf-8")
        record("illustration 'none' without a reason is rejected", any("needs a reason" in x for x in validate(deck, plan)))
        entry["illustration"] = "none: text-only grouped fixture"
        plan.write_text(json.dumps({"paper": "synthetic fixture", "reference": "fixture.pptx", "slides": [entry]}), encoding="utf-8")
        record("grouped editable text is accepted", not validate(deck, plan))
        record("reference inspection includes group children", inspect(deck)["slides"][0]["shape_count"] == 2)
        r.font.size = Pt(12)
        prs.save(deck)
        record("small text inside a group is rejected", any("16 pt" in x for x in validate(deck, plan)))
        r.font.size = Pt(20)
        image = folder / "slide-01.png"
        Image.new("RGB", (100, 100), "white").save(image)
        review = folder / "review.json"
        manifest = {"deck_sha256": sha256(deck), "slides": [{"slide": 1, "image": image.name, "image_sha256": sha256(image), "reviewed": False, "assessment": "", "issues": []}]}
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("unreviewed image is rejected", bool(check_review(deck, review)))
        manifest["slides"][0].update(reviewed=True, assessment="Synthetic checker fixture; not a real visual quality review")
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("matching complete review record is accepted", not check_review(deck, review))
        manifest["slides"][0]["auto_issues"] = ["[PowerPoint] text overlaps text: \"a\" / \"b\""]
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("unwaived automatic finding is rejected", any("neither fixed nor waived" in x for x in check_review(deck, review)))
        manifest["slides"][0]["waivers"] = {manifest["slides"][0]["auto_issues"][0]: "ok"}
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("reasonless waiver is rejected", bool(check_review(deck, review)))
        manifest["slides"][0]["waivers"] = {manifest["slides"][0]["auto_issues"][0]: "Superscript glyph bound only; render shows clear gap"}
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("specific waiver is accepted", not check_review(deck, review))
        drawn_plan = folder / "plan_drawn.json"
        drawn_plan.write_text(json.dumps({"paper": "x", "reference": "x", "slides": [dict(entry, illustration="flowchart: seed -> evolve -> evaluate loop")]}), encoding="utf-8")
        record("planned illustration that was not drawn is rejected", any("not drawn" in x for x in check_review(deck, review, drawn_plan)))
        record("'none: <reason>' plan needs no artwork", not check_review(deck, review, plan))
        reviewed_hash = manifest["deck_sha256"]
        pic = prs.slides[0].shapes.add_picture(str(image), 0, 0, 500000, 500000)
        blip = pic._element.find('.//' + qn('a:blip'))
        ext = etree.SubElement(etree.SubElement(blip, qn('a:extLst')), qn('a:ext'), uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        etree.SubElement(ext, '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip')
        prs.save(deck)
        manifest["deck_sha256"] = sha256(deck)
        review.write_text(json.dumps(manifest), encoding="utf-8")
        record("planned illustration drawn as SVG is accepted", not check_review(deck, review, drawn_plan))
        pic._element.getparent().remove(pic._element)
        manifest["deck_sha256"] = reviewed_hash
        review.write_text(json.dumps(manifest), encoding="utf-8")
        teal = ["064E41", "008380", "0E6767", "EFF7F7", "98D1CC", "0070C0"]
        moved = [shift_hue(c, 50, hue_range=(150, 215), sat_scale=0.72) for c in teal]
        record("shift_hue keeps luminance within 1%",
               all(abs(luminance(a) - luminance(b)) < 0.01 * max(luminance(a), 0.05) for a, b in zip(teal, moved)))
        record("shift_hue changes the hue of colours in range", all(a != b for a, b in zip(teal, moved)))
        record("shift_hue leaves red and greys outside the range untouched",
               shift_hue("C00000", 50, hue_range=(150, 215)) == "C00000" and shift_hue("8A8A8A", 50) == "8A8A8A")
        record("shift_svg_hue recolours artwork",
               shift_svg_hue('<path fill="#008380" stroke="#C00000"/>', 50, hue_range=(150, 215)) ==
               f'<path fill="#{shift_hue("008380", 50, hue_range=(150, 215))}" stroke="#C00000"/>')
        hp = Presentation()
        hs = hp.slides.add_slide(hp.slide_layouts[6])
        Canvas(hs.shapes, None).header(7, rule=(0, 1600, 122), backdrop=[(180, 900, 420, "EFEFF7", 0.3)])
        names = [s.name for s in hs.shapes]
        record("header draws native rule, page number and backdrop",
               names == ["底纹", "页眉线", "页码"] and hs.shapes[2].text_frame.text == "7")
        extents = {"slide_w": 960, "slide_h": 540, "slides": [{"slide": 1, "shapes": [
            {"text": "ok", "frame": [10, 10, 200, 30], "bound": [10, 12, 150, 24], "wrap": True},
            {"text": "clipped", "frame": [10, 100, 200, 30], "bound": [10, 100, 200, 80], "wrap": True},
            {"text": "collide", "frame": [12, 14, 100, 30], "bound": [12, 14, 100, 24], "wrap": False},
            {"text": "offslide", "frame": [900, 500, 200, 30], "bound": [900, 500, 150, 24], "wrap": False}]}]}
        found = " ".join(audit_layout(extents)[1])
        record("PowerPoint audit flags clipping, overlap and off-slide text",
               "taller" in found and "overlaps" in found and "outside" in found and '"ok" /' in found)
        prs.save(deck)
        record("changed PPTX invalidates old review", any("current PPTX" in x for x in check_review(deck, review)))
        prs.slides[0].shapes.add_picture(str(image), 0, 0, prs.slide_width, prs.slide_height)
        prs.save(deck)
        record("full-slide screenshot is rejected", any("full-slide picture" in x for x in validate(deck, plan)))
        plan.write_text("{}", encoding="utf-8")
        record("missing plan is rejected", bool(validate(deck, plan)))
    for name in results:
        print("PASS:", name)


if __name__ == "__main__":
    main()
