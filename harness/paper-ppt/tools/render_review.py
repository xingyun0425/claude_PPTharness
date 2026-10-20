"""Render slides in PowerPoint, audit text layout as PowerPoint measures it, and
tie a visual-review record to the exact PPTX.

    python render_review.py render deck.pptx review/          # PNGs + contact sheets + layout audit + review.json
    python render_review.py check deck.pptx review/review.json
    python render_review.py sheet refs/ [--out dir]           # contact sheets for any folder of PNG/JPG
    python render_review.py compare ref.png ours.png out.png  # side-by-side calibration image
"""
import argparse
import hashlib
import json
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from validate_deck import count_svg, illustration_plan

RENDER_W = 1600          # exported PNG width; audit coordinates use the same scale
IMAGE_EXT = ('.png', '.jpg', '.jpeg', '.webp')


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fingerprint(path):
    """Hash a file, or every image/pptx/pdf inside a reference folder."""
    path = Path(path)
    if path.is_file():
        return sha256(path)
    h = hashlib.sha256()
    for f in sorted(p for p in path.rglob('*') if p.is_file() and p.suffix.lower() in IMAGE_EXT + ('.pptx', '.pdf')):
        h.update(str(f.relative_to(path)).encode('utf-8'))
        h.update(sha256(f).encode())
    return h.hexdigest()


# ---------------------------------------------------------------- layout audit
def audit_layout(data, tol=3, tol_x=16):
    """Find clipped, overlapping or off-slide text from PowerPoint's own text extents.

    `data` is the JSON written by ppt_render.ps1 (points). Findings are reported in
    1600-px render coordinates so they match the PNGs. PowerPoint's BoundWidth runs
    ~8 px wider than the visible glyphs, hence the larger horizontal tolerance; line
    boxes include leading, so overlaps must cover a quarter of the shorter line.
    """
    scale = RENDER_W / data['slide_w']
    W, H = data['slide_w'] * scale, data['slide_h'] * scale
    found = {}
    slides = data.get('slides') or []
    for entry in slides if isinstance(slides, list) else [slides]:
        boxes, issues = [], []
        shapes = entry.get('shapes') or []
        for shp in shapes if isinstance(shapes, list) else [shapes]:
            label = shp['text'][:24]
            fx, fy, fw, fh = (v * scale for v in shp['frame'])
            bx, by, bw, bh = (v * scale for v in shp['bound'])
            if bw <= 0 or bh <= 0:
                continue
            spill_x = max(bx + bw - (fx + fw), fx - bx)
            spill_y = max(by + bh - (fy + fh), fy - by)
            if spill_x > tol_x:
                issues.append(f'text wider than its box by {spill_x:.0f}px: "{label}"')
            if spill_y > tol + 4:
                issues.append(f'text taller than its box by {spill_y:.0f}px: "{label}"')
            if bx < -tol or by < -tol or bx + bw > W + tol or by + bh > H + tol:
                issues.append(f'text outside the slide: "{label}"')
            boxes.append((bx, by, bx + bw, by + bh, label))
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                dy = min(a[3], b[3]) - max(a[1], b[1])
                if min(a[2], b[2]) - max(a[0], b[0]) > tol and dy > max(tol, 0.25 * min(a[3] - a[1], b[3] - b[1])):
                    issues.append(f'text overlaps text: "{a[4]}" / "{b[4]}"')
        found[int(entry['slide'])] = list(dict.fromkeys(issues))
    return found


# ---------------------------------------------------------------- render
def export_slides(deck, out_dir, prefix='slide', audit_json=None):
    """Run ppt_render.ps1 (local PowerPoint via COM; no pywin32 needed)."""
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File',
           str(Path(__file__).with_name('ppt_render.ps1')), '-Pptx', str(deck), '-OutDir', str(out_dir),
           '-Prefix', prefix, '-Width', str(RENDER_W)]
    if audit_json:
        cmd += ['-AuditJson', str(audit_json)]
    result = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=900)
    if result.returncode != 0:
        raise RuntimeError('PowerPoint export failed:\n' + (result.stderr or result.stdout)[-2000:])


def render(deck, out_dir, prefix='slide', audit=True, record=True, plan=None):
    """Export every slide as PNG via local PowerPoint; optionally audit layout and write review.json.

    With `plan`, each review row also records the planned illustration and how many SVG
    pieces were actually drawn on that slide.
    """
    deck = Path(deck).resolve()
    out_dir = Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob(f'{prefix}-*.png'):
        old.unlink()
    before = sha256(deck)
    audit_json = out_dir / f'{prefix}-text-extents.json' if audit else None
    export_slides(deck, out_dir, prefix, audit_json)
    found = audit_layout(json.loads(audit_json.read_text(encoding='utf-8-sig'))) if audit else {}
    count = len(list(out_dir.glob(f'{prefix}-[0-9]*.png')))
    if sha256(deck) != before:
        raise RuntimeError('PPTX changed during rendering; rerun on the saved final file')
    sheets = contact_sheets(sorted(out_dir.glob(f'{prefix}-[0-9]*.png')), out_dir / f'{prefix}-contact')
    print(f'Rendered {count} slides: {out_dir}')
    print('Contact sheets (view first for deck rhythm and consistency): ' + ', '.join(p.name for p in sheets))
    if not record:
        return count
    layout_file = deck.with_suffix('.layout.json')   # written by build_deck_template.main()
    builder_issues = {}
    if layout_file.exists():
        data = json.loads(layout_file.read_text(encoding='utf-8'))
        if data.get('deck_sha256') in (None, before):
            builder_issues = {int(k): v for k, v in data.get('slides', {}).items()}
    entries = json.loads(Path(plan).read_text(encoding='utf-8')).get('slides', []) if plan else []
    from pptx import Presentation
    deck_slides = list(Presentation(deck).slides)
    rows = []
    for n in range(1, count + 1):
        auto = [f'[PowerPoint] {x}' for x in found.get(n, [])] + [f'[builder] {x}' for x in builder_issues.get(n, [])]
        rows.append({'slide': n, 'image': f'{prefix}-{n:02d}.png',
                     'image_sha256': sha256(out_dir / f'{prefix}-{n:02d}.png'),
                     'density': density(out_dir / f'{prefix}-{n:02d}.png'),
                     'illustration_planned': (entries[n - 1].get('illustration', '') if n <= len(entries) else ''),
                     'svg_drawn': count_svg(deck_slides[n - 1]) if n <= len(deck_slides) else 0,
                     'auto_issues': auto, 'waivers': {},
                     'reviewed': False, 'assessment': '', 'issues': []})
    manifest = out_dir / 'review.json'
    manifest.write_text(json.dumps({'deck': str(deck), 'deck_sha256': before, 'slides': rows},
                                   ensure_ascii=False, indent=2), encoding='utf-8')
    flagged = [(r['slide'], x) for r in rows for x in r['auto_issues']]
    for n, x in flagged:
        print(f'  slide {n}: {x}')
    print(f'{len(flagged)} automatic layout finding(s). Fix them in the builder, or waive a specific one '
          f'in review.json "waivers" with the reason it is invisible or intended.')
    if entries:
        missing = [r['slide'] for r, e in zip(rows, entries) if illustration_plan(e)[0] and r['svg_drawn'] == 0]
        print('Illustrations (SVG pieces per slide): ' + ' '.join(f"{r['slide']}:{r['svg_drawn']}" for r in rows))
        if missing:
            print(f'Planned illustration not drawn on slides {missing}; draw it from plan.illustration before delivery.')
    print(f'Inspect every PNG, then record reviewed/assessment/issues: {manifest}')
    return count


# ---------------------------------------------------------------- images
def density(image):
    """Share of the content area (below the top 12%) that differs from the page background.

    A coarse emptiness signal for comparing a deck with its reference, not a quality score.
    """
    im = Image.open(image).convert('RGB')
    im = im.resize((320, round(im.height * 320 / im.width)))
    im = im.crop((0, round(im.height * 0.12), im.width, im.height))
    px = list(im.getdata())
    bg = max(set(px[::7]), key=px[::7].count)
    ink = sum(1 for p in px if sum(abs(a - b) for a, b in zip(p, bg)) > 40)
    return round(ink / len(px), 3)


def _font(size):
    for f in (r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\arial.ttf'):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def contact_sheets(images, stem, cols=4, per_sheet=16, thumb_w=480):
    """Numbered thumbnail grids so a whole deck can be judged at a glance."""
    images = [Path(p) for p in images]
    stem = Path(stem)
    for old in stem.parent.glob(stem.name + '-*.png'):
        old.unlink()
    out, font = [], _font(26)
    for start in range(0, len(images), per_sheet):
        chunk = images[start:start + per_sheet]
        thumbs = []
        for p in chunk:
            im = Image.open(p).convert('RGB')
            thumbs.append(im.resize((thumb_w, round(im.height * thumb_w / im.width))))
        th = max(t.height for t in thumbs)
        rows = (len(thumbs) + cols - 1) // cols
        pad, lab = 16, 34
        sheet = Image.new('RGB', (cols * (thumb_w + pad) + pad, rows * (th + lab + pad) + pad), '#D9D9D9')
        d = ImageDraw.Draw(sheet)
        for i, (p, t) in enumerate(zip(chunk, thumbs)):
            x = pad + (i % cols) * (thumb_w + pad)
            y = pad + (i // cols) * (th + lab + pad)
            d.text((x, y), p.stem, fill='#202020', font=font)
            sheet.paste(t, (x, y + lab))
        path = stem.parent / f'{stem.name}-{start // per_sheet + 1}.png'
        sheet.save(path)
        out.append(path)
    return out


def compare(reference, ours, dest, labels=('reference', 'ours')):
    """Side-by-side image at equal height for calibrating a page against its reference."""
    a, b = (Image.open(p).convert('RGB') for p in (reference, ours))
    h = 900
    a, b = (im.resize((round(im.width * h / im.height), h)) for im in (a, b))
    lab, pad = 40, 16
    canvas = Image.new('RGB', (a.width + b.width + pad * 3, h + lab + pad * 2), '#D9D9D9')
    d, font = ImageDraw.Draw(canvas), _font(28)
    for x, im, text in ((pad, a, labels[0]), (pad * 2 + a.width, b, labels[1])):
        d.text((x, pad), text, fill='#202020', font=font)
        canvas.paste(im, (x, pad + lab))
    Path(dest).parent.mkdir(parents=True, exist_ok=True)
    canvas.save(dest)
    return Path(dest)


# ---------------------------------------------------------------- review gate
def check_review(deck, manifest, plan=None):
    deck = Path(deck).resolve()
    manifest = Path(manifest).resolve()
    if not manifest.exists():
        return [f'review record missing: {manifest}']
    from pptx import Presentation
    record = json.loads(manifest.read_text(encoding='utf-8'))
    errors = []
    if record.get('deck_sha256') != sha256(deck):
        errors.append('review record does not match current PPTX; rerender and review')
    deck_slides = list(Presentation(deck).slides)
    count = len(deck_slides)
    entries = json.loads(Path(plan).read_text(encoding='utf-8')).get('slides', []) if plan else []
    slides = record.get('slides', [])
    if len(slides) != count:
        errors.append(f'review has {len(slides)} slides, deck has {count}')
    for number in range(1, count + 1):
        row = slides[number - 1] if number <= len(slides) else {}
        image = manifest.parent / row.get('image', f'slide-{number:02d}.png')
        if row.get('slide') != number:
            errors.append(f'slide {number}: review entry missing or misnumbered')
        if not image.exists() or row.get('image_sha256') != sha256(image):
            errors.append(f'slide {number}: rendered image missing or changed')
        if row.get('reviewed') is not True or row.get('issues'):
            errors.append(f'slide {number}: visual review incomplete or issues unresolved')
        if len(row.get('assessment', '').strip()) < 20:
            errors.append(f'slide {number}: specific visual assessment missing')
        if number <= len(entries) and illustration_plan(entries[number - 1])[0] and count_svg(deck_slides[number - 1]) == 0:
            errors.append(f'slide {number}: planned illustration not drawn (plan: '
                          f'"{str(entries[number - 1].get("illustration"))[:50]}"); add SVG artwork or set "none: <reason>"')
        waivers = row.get('waivers') or {}
        for issue in row.get('auto_issues', []):
            if len(str(waivers.get(issue, '')).strip()) < 10:
                errors.append(f'slide {number}: automatic finding neither fixed nor waived with a reason: {issue}')
    return errors


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest='command', required=True)
    make = sub.add_parser('render')
    make.add_argument('deck', type=Path)
    make.add_argument('output_dir', type=Path)
    check = sub.add_parser('check')
    check.add_argument('deck', type=Path)
    check.add_argument('manifest', type=Path)
    sheet = sub.add_parser('sheet')
    sheet.add_argument('folder', type=Path)
    sheet.add_argument('--cols', type=int, default=4)
    sheet.add_argument('--out', type=Path, help='output folder (default: <folder>/_sheets)')
    cmp_ = sub.add_parser('compare')
    cmp_.add_argument('reference', type=Path)
    cmp_.add_argument('ours', type=Path)
    cmp_.add_argument('dest', type=Path)
    args = parser.parse_args()
    if args.command == 'render':
        render(args.deck, args.output_dir)
    elif args.command == 'sheet':
        imgs = sorted((p for p in args.folder.iterdir() if p.suffix.lower() in IMAGE_EXT and 'contact' not in p.stem),
                      key=lambda p: (len(p.stem), p.stem))
        out = args.out or args.folder / '_sheets'
        out.mkdir(parents=True, exist_ok=True)
        for p in contact_sheets(imgs, out / 'sheet', cols=args.cols):
            print(p)
    elif args.command == 'compare':
        print(compare(args.reference, args.ours, args.dest))
    else:
        errors = check_review(args.deck, args.manifest)
        for error in errors:
            print('FAIL:', error)
        if errors:
            raise SystemExit(1)
        print('PASS: current deck has a completed visual review for every slide')


if __name__ == '__main__':
    main()
