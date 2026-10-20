"""Exercise the native builder and delivery gate with a minimal sourced slide."""
import json
import sys
import tempfile
from pathlib import Path

import build_deck_template as builder
from validate_deck import validate


def main():
    builder.THEME = {'colors': {'navy': '9A3040', 'blue': 'B24A62'}, 'font': 'Microsoft YaHei'}
    @builder.slide
    def sample(cv, number):
        cv.title('Sample finding')
        return 'Source: test fixture'

    with tempfile.TemporaryDirectory() as directory:
        deck = Path(directory) / 'sample.pptx'
        plan = Path(directory) / 'slide_plan.json'
        plan.write_text(json.dumps({
            'paper': 'test-paper.pdf',
            'reference': 'test-reference.pptx',
            'slides': [{
                'purpose': 'Show a sample finding',
                'source': 'test-paper.pdf p. 1',
                'reference_slide': 'test-reference.pptx slide 1',
                'adaptation': 'Use its simple title hierarchy',
                'visual': 'Title card; no illustration needed (smoke test fixture)',
            }],
        }), encoding='utf-8')
        sys.argv = ['build_deck_template.py', str(deck)]
        builder.main()
        from pptx import Presentation
        title = list(Presentation(deck).slides)[0].shapes[0]
        assert str(title.text_frame.paragraphs[0].runs[0].font.color.rgb) == '9A3040'
        assert validate(deck, plan) == [], 'valid native slide was rejected'
        plan.write_text('{}', encoding='utf-8')
        assert validate(deck, plan), 'missing design record was accepted'
    print('PASS: native builder and delivery gate')


if __name__ == '__main__':
    main()
