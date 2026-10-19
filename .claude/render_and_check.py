"""Render PPTX via PowerPoint COM and check for layout issues."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / 'tools'))
from deckkit import C, measure, fpx  # noqa: E402

def render_slides(pptx_path, out_dir):
    """Render all slides to JPEG via PowerPoint COM."""
    import win32com.client
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = Path(pptx_path).resolve()

    app = win32com.client.Dispatch('PowerPoint.Application')
    app.Visible = 1
    pres = app.Presentations.Open(str(pptx_path), ReadOnly=1, WithWindow=0)

    for i in range(1, pres.Slides.Count + 1):
        slide = pres.Slides(i)
        out = out_dir / f'slide-{i:02d}.jpg'
        slide.Export(str(out), 'JPG', 1600, 900)
        print(f'rendered: {out.name}')

    pres.Close()
    app.Quit()
    print(f'total: {pres.Slides.Count} slides')

def check_titles(pptx_path):
    """Check title text for potential overflow at 32pt."""
    from pptx import Presentation
    prs = Presentation(pptx_path)
    problems = []

    for i, slide in enumerate(prs.slides, 1):
        # Find title shape (name='标题')
        title_shape = None
        for sp in slide.shapes:
            if hasattr(sp, 'name') and sp.name == '标题':
                title_shape = sp
                break

        if not title_shape or not title_shape.has_text_frame:
            continue

        txt = title_shape.text_frame.text
        if not txt:
            continue

        # Check width: title box is at x=84, slides are 1600px wide, so max ~1430px
        w = measure(txt, 32, True)
        if w > 1430:
            problems.append(f'slide {i}: title "{txt}" width {w:.0f}px > 1430px')

    return problems

if __name__ == '__main__':
    pptx, out = sys.argv[1], sys.argv[2]
    render_slides(pptx, out)
    issues = check_titles(pptx)
    if issues:
        print('\ntitle width issues:')
        for p in issues:
            print(f'  {p}')
    else:
        print('\nno title width issues')
