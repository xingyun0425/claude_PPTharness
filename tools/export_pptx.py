"""Render every slide of the canvas deck with headless Chrome and pack them into a PPTX (one full-bleed image per slide).

Usage: python tools/export_pptx.py harnessdev-deck.html harnessdev-deck.pptx [--scale 2]
"""
import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
]
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9


def find_browser() -> str:
    for p in CHROME_CANDIDATES:
        if Path(p).exists():
            return p
    raise SystemExit("Chrome / Edge not found")


def slide_count(html: Path) -> int:
    return len(re.findall(r"^\s*slides\.push\(", html.read_text(encoding="utf-8"), flags=re.M))


def render(browser: str, html: Path, idx: int, out: Path, scale: float) -> None:
    url = html.resolve().as_uri() + f"?export#{idx}"
    subprocess.run(
        [browser, "--headless=new", "--disable-gpu", "--hide-scrollbars",
         "--window-size=1600,900", f"--force-device-scale-factor={scale}",
         "--virtual-time-budget=5000", f"--screenshot={out}", url],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120,
    )
    if not out.exists() or out.stat().st_size == 0:
        raise SystemExit(f"slide {idx}: screenshot failed")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("html", type=Path)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--scale", type=float, default=2.0)
    args = ap.parse_args()

    browser = find_browser()
    n = slide_count(args.html)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmp:
        for i in range(1, n + 1):
            png = Path(tmp) / f"slide-{i:02d}.png"
            render(browser, args.html, i, png, args.scale)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, SLIDE_W, SLIDE_H)
            print(f"slide {i}/{n} ok")
        prs.save(args.pptx)
    print(f"saved {args.pptx} ({args.pptx.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
