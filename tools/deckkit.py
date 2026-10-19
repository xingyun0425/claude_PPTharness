"""Toolkit for building native, editable PPTX slides in the HarnessDev deck style.

Layout coordinates are deck pixels on a 1600x900 canvas (1 px = 0.6 pt); font sizes are
PowerPoint points. Every text run goes through set_runs(), which rejects anything below MIN_PT.
"""
import re
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import quote

from lxml import etree
from PIL import Image, ImageFont
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

MIN_PT = 16
PX = 7620                      # EMU per deck pixel
PT_PER_PX = 0.6
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)
FONT = 'Microsoft YaHei'
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'
C = dict(navy='163B65', blue='347BC1', gold='D5B96F', rose='C98078', pale='EAF2FA', paleRose='F3E4E1',
         line='CFDCE8', text='243445', grey='64748B', neutral='8A99A8', white='FFFFFF', paleGold='F7F0DE',
         zone='F5F9FD', chipGrey='EEF1F5', goldLine='E9DAB3', goldText='9A7B32', darkRose='A85C54', sky='4F9AD6')
MODS = [('E', '执行', 'Execution', '163B65'), ('T', '工具', 'Tools', '347BC1'), ('C', '上下文', 'Context', '4F9AD6'),
        ('S', '状态', 'State', 'C98078'), ('L', '生命周期', 'Lifecycle', 'B39552'), ('V', '验证', 'Verification', '5B7A99')]
LINEAGE = {'Gemini': '0072B2', 'Opus': 'E69F00', 'Qwen': '009E73', 'DeepSeek': '8E44AD', 'GPT-5.5': 'D55E00'}


def E(v):
    return Emu(int(round(v * PX)))


def fpx(pt):
    """Font size in points -> deck pixels."""
    return pt / PT_PER_PX


def rgb(h):
    return RGBColor.from_string(h)


def lighten(h, t):
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return '%02X%02X%02X' % tuple(round(v + (255 - v) * t) for v in (r, g, b))


_fonts = {}
_SUB = re.compile(r'_\{([^}]*)\}|_([A-Za-z0-9])')


def plain(s):
    return _SUB.sub(lambda m: m.group(1) or m.group(2), s)


def measure(s, pt, bold=False):
    """Rendered width in deck px (subscripts measured at full size: a slight overestimate)."""
    key = (pt, bold)
    if key not in _fonts:
        path = r'C:\Windows\Fonts\msyhbd.ttc' if bold else r'C:\Windows\Fonts\msyh.ttc'
        _fonts[key] = ImageFont.truetype(path, int(round(fpx(pt) * 4)))
    # calibrated against PowerPoint's TextFrame2.BoundWidth: ~2.5% wider than PIL's advance width
    return _fonts[key].getlength(plain(s)) / 4 * 1.025


_TOKEN = re.compile(r"_\{[^}]*\}|[A-Za-z0-9.%+\-−–×/()'_:,@*†≈≥≤=]+|\s+|.")
_NO_BREAK_BEFORE = set('，。；：、）》”！？,.;:)')


def wrap_lines(s, pt, bold, width):
    """Approximate PowerPoint line breaking: CJK per character, latin per word."""
    lines, cur = [], ''
    for t in _TOKEN.findall(s):
        if measure(cur + t, pt, bold) > width * 0.97 and cur.strip() and t not in _NO_BREAK_BEFORE:
            lines.append(cur.rstrip())
            cur = t.lstrip()
        else:
            cur += t
    if cur.strip():
        lines.append(cur.rstrip())
    return lines or ['']


def set_runs(par, s, pt, bold=False, color=C['text'], italic=False):
    """Append runs to a paragraph; '_X' / '_{..}' become subscript runs."""
    if pt < MIN_PT:
        raise ValueError(f'font {pt}pt < {MIN_PT}pt: {s!r}')
    for part in [p for p in re.split(r'(_\{[^}]*\}|_[A-Za-z0-9])', s) if p]:
        sub = part.startswith('_') and len(part) > 1
        r = par.add_run()
        r.text = (part[2:-1] if part[1] == '{' else part[1]) if sub else part
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(pt), bold, italic
        f.color.rgb = rgb(color)
        rPr = r._r.get_or_add_rPr()
        rPr.set('lang', 'zh-CN')
        for tag in ('a:ea', 'a:cs'):
            etree.SubElement(rPr, qn(tag), typeface=FONT)
        if sub:
            rPr.set('baseline', '-25000')


def frame(tf, anchor=MSO_ANCHOR.MIDDLE, wrap=False):
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor


def strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, alpha=14, blur=16, dist=5):
    eff = etree.SubElement(shape._element.spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad=str(blur * PX), dist=str(dist * PX),
                          dir='5400000', algn='t', rotWithShape='0')
    clr = etree.SubElement(sh, qn('a:srgbClr'), val=C['navy'])
    etree.SubElement(clr, qn('a:alpha'), val=str(alpha * 1000))


ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
class Canvas:
    """Drawing API over a python-pptx shape collection, in deck pixels."""

    def __init__(self, shapes, icons, log=None):
        self.shapes, self.icons = shapes, icons
        self.log = [] if log is None else log   # [(x0, y0, x1, y1, label)] of rendered text extents, shared by groups
        self.overflow = []
        self.cards = []   # bounding boxes of cards/zones that text is expected to stay inside

    def group(self, name):
        g = self.shapes.add_group_shape()
        g.name = name
        sub = Canvas(g.shapes, self.icons, self.log)
        sub.overflow, sub.cards = self.overflow, self.cards
        return sub

    def _reg(self, x0, y0, x1, y1, label):
        self.log.append((x0, y0, x1, y1, label))

    def overlaps(self, tol=2):
        """Pairs of text extents that collide, text that starts inside a card but spills out, and off-slide text."""
        out = []
        for t in self.log:
            for (bx0, by0, bx1, by1, bname) in self.cards:
                cx, cy = t[0] + 4, (t[1] + t[3]) / 2
                if bx0 < cx < bx1 and by0 < cy < by1 and (t[2] > bx1 + tol or t[0] < bx0 - tol):
                    out.append(('溢出卡片', t[4], bname))
        for i, a in enumerate(self.log):
            if a[0] < -tol or a[1] < -tol or a[2] > 1600 + tol or a[3] > 900 + tol:
                out.append(('出界', a[4], ''))
            for b in self.log[i + 1:]:
                if min(a[2], b[2]) - max(a[0], b[0]) > tol and min(a[3], b[3]) - max(a[1], b[1]) > tol:
                    out.append(('重叠', a[4], b[4]))
        return out + [('溢出', o, '') for o in self.overflow]

    # ---- text ----
    def text(self, s, x, y, pt=18, bold=False, color=C['text'], align='l', name=None):
        """Single line; (x, y) = anchor point at vertical centre (left/centre/right by align)."""
        tw = measure(s, pt, bold)
        w = tw * 1.05 + 8
        h = fpx(pt) * 1.5
        left = {'l': x, 'c': x - w / 2, 'r': x - w}[align]
        tb = self.shapes.add_textbox(E(left), E(y - h / 2), E(w), E(h))
        frame(tb.text_frame)
        p = tb.text_frame.paragraphs[0]
        p.alignment = ALIGN[align]
        set_runs(p, s, pt, bold, color)
        if name:
            tb.name = name
        tl = {'l': x, 'c': x - tw / 2, 'r': x - tw}[align]
        self._reg(tl, y - fpx(pt) * 0.5, tl + tw, y + fpx(pt) * 0.5, plain(s))
        return tb

    def rich(self, segs, x, y, pt=18, align='l', name=None):
        """One line built from [(text, {bold, color}), ...]."""
        tw = sum(measure(t, pt, o.get('bold', False)) for t, o in segs)
        w = tw * 1.05 + 8
        h = fpx(pt) * 1.5
        left = {'l': x, 'c': x - w / 2, 'r': x - w}[align]
        tl = {'l': x, 'c': x - tw / 2, 'r': x - tw}[align]
        self._reg(tl, y - fpx(pt) * 0.5, tl + tw, y + fpx(pt) * 0.5, ''.join(plain(t) for t, _ in segs))
        tb = self.shapes.add_textbox(E(left), E(y - h / 2), E(w), E(h))
        frame(tb.text_frame)
        p = tb.text_frame.paragraphs[0]
        p.alignment = ALIGN[align]
        for t, o in segs:
            set_runs(p, t, pt, o.get('bold', False), o.get('color', C['text']))
        if name:
            tb.name = name
        return tb

    def para(self, s, x, y, w, pt=18, bold=False, color=C['text'], lh=1.3, align='l', name=None):
        """Wrapped paragraph with top-left at (x, y); returns bottom y (estimated)."""
        lines = wrap_lines(s, pt, bold, w)
        n = len(lines)
        h = n * fpx(pt) * lh * 1.2
        widest = max(measure(l, pt, bold) for l in lines)
        self._reg(x, y, x + widest, y + h, plain(s)[:16] + '…')
        tb = self.shapes.add_textbox(E(x), E(y), E(w), E(h))
        frame(tb.text_frame, anchor=MSO_ANCHOR.TOP, wrap=True)
        p = tb.text_frame.paragraphs[0]
        p.alignment = ALIGN[align]
        p.line_spacing = lh
        set_runs(p, s, pt, bold, color)
        if name:
            tb.name = name
        return y + h

    def bullets(self, items, x, y, w, pt=18, color=C['text'], dot=C['blue'], gap=12, lh=1.3):
        for s in items:
            self.box(x, y + fpx(pt) * 0.55, 10, 10, fill=dot, name='项目符号')
            y = self.para(s, x + 24, y, w - 24, pt=pt, color=color, lh=lh) + gap
        return y

    # ---- shapes ----
    def box(self, x, y, w, h, fill=None, line=None, lw=2, r=0, dash=False, shadow=False, shape=None, name=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
        sp = self.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
        strip_style(sp)
        if r and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r / max(1, min(w, h)))
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Emu(int(lw * PX))
            if dash:
                etree.SubElement(sp.line._get_or_add_ln(), qn('a:prstDash'), val='dash')
        else:
            sp.line.fill.background()
        if shadow:
            add_shadow(sp)
        frame(sp.text_frame)
        if name:
            sp.name = name
        return sp

    def label_box(self, s, x, y, w, h, pt=18, bold=True, color=C['white'], align='c', **kw):
        tw = measure(s, pt, bold)
        if tw > w - 8:
            self.overflow.append(f'[溢出框] {plain(s)}')
        cx = x + w / 2
        self._reg(cx - min(tw, w) / 2, y + h / 2 - fpx(pt) * 0.5, cx + min(tw, w) / 2, y + h / 2 + fpx(pt) * 0.5, plain(s))
        sp = self.box(x, y, w, h, **kw)
        p = sp.text_frame.paragraphs[0]
        p.alignment = ALIGN[align]
        set_runs(p, s, pt, bold, color)
        return sp

    def pill(self, s, cx, cy, pt=16, fill=C['navy'], color=C['white'], line=None, padX=16, h=None, bold=True, shadow=False):
        h = h or fpx(pt) * 1.75
        w = measure(s, pt, bold) + padX * 2
        self._reg(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2, f'标签 {plain(s)}')
        sp = self.box(cx - w / 2, cy - h / 2, w, h, fill=fill, line=line, r=h / 2, shadow=shadow, name=f'标签 {plain(s)}')
        p = sp.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_runs(p, s, pt, bold, color)
        return sp

    @staticmethod
    def pill_width(s, pt=16, padX=16, bold=True):
        return measure(s, pt, bold) + padX * 2

    def pill_row(self, labels, cx, cy, gap=8, **kw):
        pt, padX, bold = kw.get('pt', 16), kw.get('padX', 16), kw.get('bold', True)
        ws = [self.pill_width(s, pt, padX, bold) for s in labels]
        x = cx - (sum(ws) + gap * (len(labels) - 1)) / 2
        for s, w in zip(labels, ws):
            self.pill(s, x + w / 2, cy, **kw)
            x += w + gap

    def circle(self, cx, cy, r, fill=None, line=None, lw=3, shadow=False, name='圆'):
        return self.box(cx - r, cy - r, 2 * r, 2 * r, fill=fill, line=line, lw=lw, shadow=shadow, shape=MSO_SHAPE.OVAL, name=name)

    def num(self, n, cx, cy, r=22, fill=C['blue'], pt=18):
        sp = self.circle(cx, cy, r, fill=fill, name=f'序号 {n}')
        p = sp.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_runs(p, str(n), pt, True, C['white'])
        return sp

    def halo(self, cx, cy, r, fill=C['pale']):
        return self.circle(cx, cy, r, fill=fill, shadow=True, name='底托')

    def line(self, x1, y1, x2, y2, color=C['line'], lw=2, dash=False, name='线'):
        sp = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        strip_style(sp)
        sp.line.color.rgb = rgb(color)
        sp.line.width = Emu(int(lw * PX))
        if dash:
            etree.SubElement(sp.line._get_or_add_ln(), qn('a:prstDash'), val='dash')
        sp.name = name
        return sp

    def arrow(self, x1, y1, x2, y2, color=C['blue'], lw=3.5, dash=False, ctrl=None, both=False):
        sp = self.line(x1, y1, x2, y2, color, lw, dash, '箭头') if ctrl is None else self._quad(x1, y1, ctrl, x2, y2, color, lw, dash)
        ln = sp.line._get_or_add_ln()
        etree.SubElement(ln, qn('a:round'))
        if both:
            etree.SubElement(ln, qn('a:headEnd'), type='triangle', w='med', len='med')
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
        return sp

    def _quad(self, x1, y1, c, x2, y2, color, lw, dash):
        xs, ys = (x1, c[0], x2), (y1, c[1], y2)
        L, T = min(xs), min(ys)
        W, H = max(max(xs) - L, 1), max(max(ys) - T, 1)
        sp = self.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(L), E(T), E(W), E(H))
        strip_style(sp)
        sp.fill.background()
        spPr = sp._element.spPr
        spPr.remove(spPr.find(qn('a:prstGeom')))
        pt = lambda x, y: f'<a:pt x="{int((x - L) * PX)}" y="{int((y - T) * PX)}"/>'
        geom = etree.fromstring(
            f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
            f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst><a:path w="{int(W * PX)}" h="{int(H * PX)}" fill="none">'
            f'<a:moveTo>{pt(x1, y1)}</a:moveTo><a:quadBezTo>{pt(*c)}{pt(x2, y2)}</a:quadBezTo></a:path></a:pathLst></a:custGeom>')
        spPr.insert(1, geom)
        sp.line.color.rgb = rgb(color)
        sp.line.width = Emu(int(lw * PX))
        if dash:
            etree.SubElement(sp.line._get_or_add_ln(), qn('a:prstDash'), val='dash')
        sp.name = '弧形箭头'
        return sp

    def block_arrow(self, x, y, w, h, fill=C['gold'], name='过渡箭头'):
        a = self.box(x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW, name=name)
        a.adjustments[0], a.adjustments[1] = 0.5, 0.45
        return a

    def icon(self, key, x, y, w, h=None, name=None):
        return self.icons.place(self.shapes, key, x, y, w, h, name or key)

    def icon_c(self, key, cx, cy, size, name=None):
        """Place an icon fitted into a size x size square centred at (cx, cy)."""
        vw, vh = self.icons.aspect(key)
        w, h = (size, size * vh / vw) if vw >= vh else (size * vw / vh, size)
        return self.icon(key, cx - w / 2, cy - h / 2, w, h, name)
    # ---- deck chrome ----
    def title(self, s):
        self.text(s, 84, 78, pt=32, bold=True, color=C['navy'], name='标题')

    def page(self, n):
        self.text(str(n), 1516, 850, pt=16, color=C['grey'], align='r', name='页码')

    def band(self, s, y=752, h=64, x=84, w=1432, pt=20):
        self.label_box(s, x, y, w, h, pt=pt, bold=True, color=C['navy'], fill=C['pale'], name='结论条')

    def panel(self, x, y, w, h, head, fill=C['navy'], head_h=56, pt=19):
        """White card with a solid header bar (deck's standard panel)."""
        self.cards.append((x, y + head_h, x + w, y + h, f'面板 {head}'))
        self.box(x, y + head_h, w, h - head_h, fill=C['white'], line=C['line'], name=f'面板 {head}')
        self.label_box(head, x, y, w, head_h, pt=pt, fill=fill, name=f'面板标题 {head}')

    def zone(self, x, y, w, h, head=None, head_fill=C['navy'], pt=19):
        """Rounded light zone with a floating pill heading (illustration style)."""
        self.cards.append((x, y + 24, x + w, y + h, f'分区 {head or ""}'))
        self.box(x, y, w, h, fill=C['zone'], line=C['line'], r=18, name=f'分区 {head or ""}')
        if head:
            self.pill(head, x + w / 2, y, pt=pt, fill=head_fill, padX=26, shadow=True)

    def card(self, x, y, w, h, fill=C['white'], line=C['line'], r=14, shadow=True, name='卡片'):
        self.cards.append((x, y, x + w, y + h, name))
        return self.box(x, y, w, h, fill=fill, line=line, r=r, shadow=shadow, name=name)

    def kpi(self, x, y, w, h, value, label, dark=False, vpt=32, lpt=16):
        self.card(x, y, w, h, fill=C['navy'] if dark else C['pale'], line=None, shadow=dark, name=f'指标 {value}')
        self.text(value, x + w / 2, y + h * 0.40, pt=vpt, bold=True, color=C['white'] if dark else C['navy'], align='c')
        self.text(label, x + w / 2, y + h * 0.78, pt=lpt, color=C['pale'] if dark else C['grey'], align='c')

    def mod_tile(self, k, zh, col, x, y, w=60, h=60, kpt=20, zpt=None):
        tile = self.box(x, y, w, h, fill=col, r=10, name=f'模块 {k}')
        tf = tile.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.9
        set_runs(p, k, kpt, True, C['white'])
        if zpt:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.line_spacing = 0.9
            set_runs(p2, zh, zpt, True, C['white'])
        return tile

    def table(self, x, y, col_w, header, rows, row_h=46, head_h=50, pt=16, head_pt=16,
              head_fills=None, row_fills=None, styles=None):
        """Native PowerPoint table. styles(i, j) -> dict(bold, color) overrides per body cell."""
        gf = self.shapes.add_table(len(rows) + 1, len(col_w), E(x), E(y), E(sum(col_w)), E(head_h + row_h * len(rows)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for attr in ('firstRow', 'bandRow'):
            tblPr.set(attr, '0')
        style_id = tblPr.find(qn('a:tableStyleId'))
        if style_id is not None:
            style_id.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # "No Style, Table Grid"
        for j, w in enumerate(col_w):
            tbl.columns[j].width = E(w)
        tbl.rows[0].height = E(head_h)
        for i in range(len(rows)):
            tbl.rows[i + 1].height = E(row_h)

        def fill_cell(cell, s, size, bold, color, bg):
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg)
            cell.margin_left = cell.margin_right = E(8)
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_runs(p, s, size, bold, color)
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(side), w=str(int(1.5 * PX)))
                if side == 'a:lnB':
                    sf = etree.SubElement(ln, qn('a:solidFill'))
                    etree.SubElement(sf, qn('a:srgbClr'), val=C['line'])
                else:
                    etree.SubElement(ln, qn('a:noFill'))

        for j, h in enumerate(header):
            hf = (head_fills or [C['navy']] * len(header))[j]
            fill_cell(tbl.cell(0, j), h, head_pt, True, C['white'], hf)
        for i, r in enumerate(rows):
            bg = (row_fills[i] if row_fills else (C['pale'] if i % 2 == 0 else C['white']))
            for j, v in enumerate(r):
                st = (styles(i, j) if styles else None) or {}
                fill_cell(tbl.cell(i + 1, j), str(v), st.get('pt', pt), st.get('bold', j == 0), st.get('color', C['navy'] if j == 0 else C['text']), st.get('bg', bg))
        gf.name = '表格'
        return gf

    def line_chart(self, x, y, w, h, categories, series, y_max, y_min=0, major=None, legend=False, title=None, pt=16):
        """Native line chart; series = [(name, values, hex color, dashed)]. None values are gaps."""
        cd = CategoryChartData()
        cd.categories = categories
        for name, vals, _, _ in series:
            cd.add_series(name, vals)
        gf = self.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, E(x), E(y), E(w), E(h), cd)
        ch = gf.chart
        ch.font.name, ch.font.size = FONT, Pt(pt)
        ch.font.color.rgb = rgb(C['grey'])
        ch.has_legend = legend
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            r = ch.chart_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(18), True, FONT
            r.font.color.rgb = rgb(C['navy'])
        else:
            ch.has_title = False
        va = ch.value_axis
        va.maximum_scale, va.minimum_scale = y_max, y_min
        if major:
            va.major_unit = major
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(C['line'])
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(pt)
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(pt)
        ca.format.line.color.rgb = rgb(C['grey'])
        for s, (_, _, col, dashed) in zip(ch.plots[0].series, series):
            s.smooth = False
            s.format.line.color.rgb = rgb(col)
            s.format.line.width = Pt(2.25)
            if dashed:
                s.format.line.dash_style = 4  # dash
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 5
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(col)
            s.marker.format.line.fill.background()
        # show gaps for missing points
        plot = ch._chartSpace.find('.//' + qn('c:chart'))
        dba = plot.find(qn('c:dispBlanksAs'))
        if dba is None:
            dba = etree.SubElement(plot, qn('c:dispBlanksAs'))
        dba.set('val', 'gap')
        gf.name = f'折线图 {title or ""}'
        return gf


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s
J = 'stroke-linejoin="round" stroke-linecap="round"'
NAVY, WHITE = '#' + C['navy'], '#FFFFFF'


def _svg(vb, body):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{vb}">{body}</svg>'


def svg_robot(accent):
    a = '#' + accent
    return _svg('0 0 100 120',
                f'<g stroke="{NAVY}" stroke-width="3.2" {J}>'
                f'<line x1="50" y1="10" x2="50" y2="24"/><circle cx="50" cy="8" r="6" fill="{a}"/>'
                f'<rect x="6" y="38" width="10" height="24" rx="4" fill="{a}"/><rect x="84" y="38" width="10" height="24" rx="4" fill="{a}"/>'
                f'<rect x="13" y="24" width="74" height="56" rx="19" fill="{WHITE}"/>'
                f'<path d="M18,118 Q18,86 50,86 Q82,86 82,118 Z" fill="{WHITE}"/></g>'
                f'<rect x="24" y="35" width="52" height="32" rx="12" fill="{NAVY}"/>'
                f'<circle cx="38" cy="51" r="5.5" fill="#9FD4FF"/><circle cx="62" cy="51" r="5.5" fill="#9FD4FF"/>')


def svg_doc(stroke, line_color, dashed=False):
    s, lc = '#' + stroke, '#' + line_color
    d = ' stroke-dasharray="7 6"' if dashed else ''
    dl = ' stroke-dasharray="6 7"' if dashed else ''
    return _svg('0 0 80 100',
                f'<g fill="none" stroke-width="3.2" {J}><path d="M4,4 H58 L76,22 V96 H4 Z" fill="{WHITE}" stroke="{s}"{d}/>'
                f'<path d="M58,4 V22 H76" stroke="{s}"/>'
                + ''.join(f'<line x1="16" y1="{y}" x2="{x2}" y2="{y}" stroke="{lc}" stroke-width="4.5"{dl}/>' for y, x2 in ((44, 64), (62, 64), (80, 42)))
                + '</g>')


def svg_lock(color):
    c = '#' + color
    return _svg('0 0 30 34', f'<g stroke="{c}" stroke-width="3" {J}><path d="M7,15 V10 A8,8 0 0 1 23,10 V15" fill="none"/>'
                             f'<rect x="3" y="14" width="24" height="18" rx="4" fill="{WHITE}"/></g><circle cx="15" cy="23" r="2.6" fill="{c}"/>')


def svg_eye(color):
    c = '#' + color
    return _svg('0 0 36 22', f'<path d="M2,11 Q18,-3 34,11 Q18,25 2,11 Z" fill="{WHITE}" stroke="{c}" stroke-width="3" {J}/>'
                             f'<circle cx="18" cy="11" r="5" fill="{c}"/>')


def svg_check(color):
    c = '#' + color
    return _svg('0 0 34 34', f'<circle cx="17" cy="17" r="15" fill="{c}" stroke="{WHITE}" stroke-width="3"/>'
                             f'<path d="M10,17 L15,22 L24,12" fill="none" stroke="{WHITE}" stroke-width="3.6" {J}/>')


def svg_cross(color):
    c = '#' + color
    return _svg('0 0 34 34', f'<circle cx="17" cy="17" r="15" fill="{c}" stroke="{WHITE}" stroke-width="3"/>'
                             f'<path d="M11,11 L23,23 M23,11 L11,23" stroke="{WHITE}" stroke-width="3.6" {J}/>')


def svg_warn(color):
    c = '#' + color
    return _svg('0 0 40 36', f'<path d="M20,3 L37,33 H3 Z" fill="{c}" stroke="{c}" stroke-width="3" {J}/>'
                             f'<line x1="20" y1="13" x2="20" y2="22" stroke="{WHITE}" stroke-width="3.6" {J}/><circle cx="20" cy="28" r="2.2" fill="{WHITE}"/>')


def svg_terminal():
    b = '#' + C['blue']
    return _svg('0 0 96 76', f'<rect x="3" y="3" width="90" height="70" rx="10" fill="{WHITE}" stroke="{NAVY}" stroke-width="3.2"/>'
                             f'<path d="M3,20 H93" stroke="{NAVY}" stroke-width="3"/><circle cx="14" cy="12" r="3" fill="#C98078"/>'
                             f'<circle cx="24" cy="12" r="3" fill="#D5B96F"/><circle cx="34" cy="12" r="3" fill="{b}"/>'
                             f'<path d="M18,36 L30,46 L18,56" fill="none" stroke="{b}" stroke-width="4" {J}/>'
                             f'<line x1="38" y1="56" x2="62" y2="56" stroke="{b}" stroke-width="4" {J}/>')


def svg_code():
    b = '#' + C['blue']
    return _svg('0 0 80 100', f'<g fill="none" stroke-width="3.2" {J}><path d="M4,4 H58 L76,22 V96 H4 Z" fill="{WHITE}" stroke="{NAVY}"/>'
                              f'<path d="M58,4 V22 H76" stroke="{NAVY}"/><path d="M28,46 L16,58 L28,70" stroke="{b}" stroke-width="4.5"/>'
                              f'<path d="M52,46 L64,58 L52,70" stroke="{b}" stroke-width="4.5"/><path d="M44,42 L36,74" stroke="#D5B96F" stroke-width="4.5"/></g>')


def svg_chart():
    return _svg('0 0 90 80', f'<rect x="3" y="3" width="84" height="74" rx="10" fill="{WHITE}" stroke="{NAVY}" stroke-width="3.2"/>'
                             f'<rect x="16" y="44" width="12" height="22" rx="2" fill="#8A99A8"/><rect x="38" y="30" width="12" height="36" rx="2" fill="#D5B96F"/>'
                             f'<rect x="60" y="16" width="12" height="50" rx="2" fill="#347BC1"/>')


def svg_pen():
    return _svg('0 0 80 100', f'<g fill="none" stroke-width="3.2" {J}><path d="M4,4 H58 L76,22 V96 H4 Z" fill="{WHITE}" stroke="{NAVY}"/>'
                              f'<path d="M58,4 V22 H76" stroke="{NAVY}"/><line x1="16" y1="44" x2="54" y2="44" stroke="#347BC1" stroke-width="4.5"/>'
                              f'<line x1="16" y1="60" x2="40" y2="60" stroke="#347BC1" stroke-width="4.5"/></g>'
                              f'<path d="M46,86 L50,72 L72,50 L80,58 L58,80 Z" fill="#D5B96F" stroke="{NAVY}" stroke-width="3" {J}/>')


def svg_search():
    return _svg('0 0 80 80', f'<circle cx="34" cy="34" r="24" fill="{WHITE}" stroke="{NAVY}" stroke-width="4"/>'
                             f'<circle cx="34" cy="34" r="14" fill="#EAF2FA"/><line x1="52" y1="52" x2="74" y2="74" stroke="{NAVY}" stroke-width="8" {J}/>')


def svg_gear(color):
    import math
    c = '#' + color
    pts = []
    for i in range(16):
        r = 34 if i % 2 == 0 else 26
        a = math.pi * 2 * i / 16
        pts.append(f'{40 + r * math.cos(a):.1f},{40 + r * math.sin(a):.1f}')
    return _svg('0 0 80 80', f'<polygon points="{" ".join(pts)}" fill="{c}" stroke="{c}" stroke-width="6" {J}/>'
                             f'<circle cx="40" cy="40" r="12" fill="{WHITE}"/>')


def svg_loop():
    b = '#' + C['blue']
    return _svg('0 0 80 80', f'<path d="M64,30 A26,26 0 1 0 66,46" fill="none" stroke="{b}" stroke-width="6" {J}/>'
                             f'<path d="M56,22 L66,30 L56,38" fill="none" stroke="{b}" stroke-width="6" {J}/>')


def svg_funnel():
    return _svg('0 0 80 84', f'<path d="M4,6 H76 L48,42 V74 L32,80 V42 Z" fill="#EAF2FA" stroke="{NAVY}" stroke-width="3.4" {J}/>'
                             f'<circle cx="24" cy="18" r="4" fill="#D5B96F"/><circle cx="40" cy="20" r="4" fill="#347BC1"/><circle cx="56" cy="17" r="4" fill="#C98078"/>')


def svg_folder():
    return _svg('0 0 96 76', f'<path d="M4,12 Q4,6 10,6 H36 L44,16 H86 Q92,16 92,22 V66 Q92,72 86,72 H10 Q4,72 4,66 Z" fill="#EAF2FA" stroke="{NAVY}" stroke-width="3.2" {J}/>'
                             f'<path d="M4,28 H92" stroke="{NAVY}" stroke-width="3"/>')


def svg_shield():
    return _svg('0 0 70 80', f'<path d="M35,4 L64,14 V38 Q64,62 35,76 Q6,62 6,38 V14 Z" fill="#EAF2FA" stroke="{NAVY}" stroke-width="3.4" {J}/>'
                             f'<path d="M22,40 L32,50 L50,30" fill="none" stroke="#347BC1" stroke-width="5" {J}/>')


def svg_flag():
    return _svg('0 0 70 80', f'<line x1="12" y1="6" x2="12" y2="76" stroke="{NAVY}" stroke-width="4" {J}/>'
                             f'<path d="M12,8 H60 L50,24 L60,40 H12 Z" fill="#D5B96F" stroke="{NAVY}" stroke-width="3" {J}/>')


def svg_star(color):
    import math
    c = '#' + color
    pts = ' '.join(f'{20 + (18 if i % 2 == 0 else 8) * math.cos(-math.pi / 2 + i * math.pi / 5):.1f},'
                   f'{20 + (18 if i % 2 == 0 else 8) * math.sin(-math.pi / 2 + i * math.pi / 5):.1f}' for i in range(10))
    return _svg('0 0 40 40', f'<polygon points="{pts}" fill="{c}" stroke="{WHITE}" stroke-width="2" {J}/>')


def svg_bulb():
    return _svg('0 0 60 80', f'<path d="M30,6 A22,22 0 0 1 44,45 Q40,50 40,58 H20 Q20,50 16,45 A22,22 0 0 1 30,6 Z" fill="#F7F0DE" stroke="{NAVY}" stroke-width="3.2" {J}/>'
                             f'<line x1="21" y1="66" x2="39" y2="66" stroke="{NAVY}" stroke-width="3.2" {J}/>'
                             f'<line x1="24" y1="74" x2="36" y2="74" stroke="{NAVY}" stroke-width="3.2" {J}/>'
                             f'<path d="M24,34 L30,42 L36,34" fill="none" stroke="#D5B96F" stroke-width="3" {J}/>')


ICON_SVG = {
    'robot_blue': svg_robot(C['blue']), 'robot_grey': svg_robot(C['neutral']), 'robot_gold': svg_robot(C['gold']),
    'doc_grey': svg_doc(C['neutral'], C['line']), 'doc_seed': svg_doc(C['rose'], C['rose'], True), 'doc_blue': svg_doc(C['navy'], C['blue']),
    'lock_navy': svg_lock(C['navy']), 'lock_grey': svg_lock(C['neutral']), 'eye_blue': svg_eye(C['blue']),
    'check_blue': svg_check(C['blue']), 'cross_rose': svg_cross(C['rose']), 'warn_rose': svg_warn(C['rose']),
    'terminal': svg_terminal(), 'code': svg_code(), 'chart': svg_chart(), 'pen': svg_pen(), 'search': svg_search(),
    'gear_blue': svg_gear(C['blue']), 'gear_grey': svg_gear(C['neutral']), 'loop': svg_loop(), 'funnel': svg_funnel(),
    'folder': svg_folder(), 'shield': svg_shield(), 'flag': svg_flag(), 'star_navy': svg_star(C['navy']),
    'star_gold': svg_star(C['gold']), 'bulb': svg_bulb(),
}
ICON_LABEL = {
    'robot_blue': '机器人·蓝', 'robot_grey': '机器人·灰', 'robot_gold': '机器人·金', 'doc_grey': '文档·灰', 'doc_seed': '种子文档',
    'doc_blue': '文档·蓝', 'lock_navy': '锁·深蓝', 'lock_grey': '锁·灰', 'eye_blue': '可见', 'check_blue': '对勾', 'cross_rose': '叉号',
    'warn_rose': '警示', 'terminal': '终端', 'code': '代码', 'chart': '图表', 'pen': '写作', 'search': '检索', 'gear_blue': '齿轮·蓝',
    'gear_grey': '齿轮·灰', 'loop': '循环', 'funnel': '漏斗', 'folder': '文件夹', 'shield': '盾牌', 'flag': '旗帜',
    'star_navy': '星标·深蓝', 'star_gold': '星标·金', 'bulb': '灯泡',
}
class IconStore:
    """Embeds each icon once as image/svg+xml (asvg:svgBlip) with a Chrome-rendered PNG fallback."""

    def __init__(self, prs, tmp):
        self.prs, self.tmp, self.png, self.parts = prs, Path(tmp), {}, {}

    def aspect(self, key):
        vb = [float(v) for v in re.search(r'viewBox="([^"]+)"', ICON_SVG[key]).group(1).split()]
        return vb[2], vb[3]

    def render_all(self):
        """Render all PNG fallbacks in one Chrome call (a sprite sheet), then crop."""
        keys = list(ICON_SVG)
        cell = 640
        cols = 6
        rows = (len(keys) + cols - 1) // cols
        imgs = []
        for i, k in enumerate(keys):
            vw, vh = self.aspect(k)
            s = 600 / max(vw, vh)
            w, h = int(vw * s), int(vh * s)
            x, y = (i % cols) * cell, (i // cols) * cell
            imgs.append((k, x, y, w, h))
        html = ''.join(f'<img src="data:image/svg+xml;charset=utf-8,{quote(ICON_SVG[k])}" width="{w}" height="{h}" '
                       f'style="position:absolute;left:{x}px;top:{y}px">' for k, x, y, w, h in imgs)
        page = self.tmp / 'sprite.html'
        page.write_text(f'<html><body style="margin:0;background:transparent">{html}</body></html>', encoding='utf-8')
        sheet = self.tmp / 'sprite.png'
        subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--default-background-color=00000000',
                        f'--window-size={cols * cell},{rows * cell}', '--force-device-scale-factor=1',
                        f'--screenshot={sheet}', page.as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        im = Image.open(sheet).convert('RGBA')
        for k, x, y, w, h in imgs:
            out = self.tmp / f'{k}.png'
            im.crop((x, y, x + w, y + h)).save(out)
            self.png[k] = out

    def place(self, shapes, key, x, y, w, h, name):
        if not self.png:
            self.render_all()
        if h is None:
            vw, vh = self.aspect(key)
            h = w * vh / vw
        pic = shapes.add_picture(str(self.png[key]), E(x), E(y), E(w), E(h))
        pic.name = name if name != key else ICON_LABEL.get(key, key)
        if key not in self.parts:
            self.parts[key] = Part(PackURI(f'/ppt/media/icon-{key}.svg'), 'image/svg+xml', self.prs.part.package,
                                   ICON_SVG[key].encode('utf-8'))
        rId = shapes.part.relate_to(self.parts[key], RT.IMAGE)
        blip = pic._element.find('.//' + qn('a:blip'))
        ext = etree.SubElement(etree.SubElement(blip, qn('a:extLst')), qn('a:ext'), uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        svgb = etree.SubElement(ext, '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
                                nsmap={'asvg': 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
        svgb.set(qn('r:embed'), rId)
        return pic
