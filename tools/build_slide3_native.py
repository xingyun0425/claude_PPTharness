"""Build slide 3 of the HarnessDev deck directly as a PPTX.

Text, cards, pills, arrows and module tiles are native PowerPoint shapes (editable in place);
only the pictorial icons are embedded SVG (with PNG fallback), which PowerPoint can recolor or
"Convert to Shape". Coordinates are written in the deck's 1600x900 px space.

Usage: python -I tools/build_slide3_native.py out.pptx
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import quote

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

C = dict(navy='163B65', blue='347BC1', gold='D5B96F', rose='C98078', pale='EAF2FA', paleRose='F3E4E1',
         line='CFDCE8', text='243445', grey='64748B', neutral='8A99A8', white='FFFFFF', paleGold='F7F0DE',
         zone='F5F9FD', chipGrey='EEF1F5', goldLine='E9DAB3', goldText='A88A3F')
MODS = [('E', '执行', '163B65'), ('T', '工具', '347BC1'), ('C', '上下文', '4F9AD6'),
        ('S', '状态', 'C98078'), ('L', '生命周期', 'B39552'), ('V', '验证', '5B7A99')]
FONT = 'Microsoft YaHei'
PX = 7620                      # EMU per deck pixel (12192000 / 1600)
PT_PER_PX = 0.6                # 960 pt slide width / 1600 px
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'
_fonts = {}


def E(v):
    return Emu(int(round(v * PX)))


def measure(s, size, bold):
    key = (size, bold)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(r'C:\Windows\Fonts\msyhbd.ttc' if bold else r'C:\Windows\Fonts\msyh.ttc', int(round(size * 4)))
    plain = re.sub(r'_\{([^}]*)\}|_(.)', lambda m: m.group(1) or m.group(2), s)
    return _fonts[key].getlength(plain) / 4


def rgb(h):
    return RGBColor.from_string(h)


def strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, alpha=13, blur=14, dist=5):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad=str(blur * PX), dist=str(dist * PX), dir='5400000', algn='t', rotWithShape='0')
    clr = etree.SubElement(sh, qn('a:srgbClr'), val=C['navy'])
    etree.SubElement(clr, qn('a:alpha'), val=str(alpha * 1000))
def set_runs(tf_par, s, size, bold, color):
    """Fill a paragraph with runs; '_X' / '_{..}' become subscript runs."""
    for part in [p for p in re.split(r'(_\{[^}]*\}|_[A-Za-z0-9])', s) if p]:
        sub = part.startswith('_')
        r = tf_par.add_run()
        r.text = (part[2:-1] if part[1] == '{' else part[1]) if sub else part
        f = r.font
        f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size * PT_PER_PX), bold, rgb(color)
        rPr = r._r.get_or_add_rPr()
        rPr.set('lang', 'zh-CN')
        for tag in ('a:ea', 'a:cs'):
            etree.SubElement(rPr, qn(tag), typeface=FONT)
        if sub:
            rPr.set('baseline', '-25000')


def frame(tf, anchor=MSO_ANCHOR.MIDDLE, wrap=False):
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor


class Canvas:
    """Thin drawing API over a python-pptx shape collection, in deck pixels."""

    def __init__(self, shapes, svg_store):
        self.shapes, self.svg = shapes, svg_store

    def group(self, name):
        g = self.shapes.add_group_shape()
        g.name = name
        return Canvas(g.shapes, self.svg)

    def text(self, s, x, y, size=20, bold=False, color=C['text'], anchor='middle', name=None):
        w = measure(s, size, bold) * 1.06 + 6
        h = size * 1.5
        left = x - w / 2 if anchor == 'middle' else x
        tb = self.shapes.add_textbox(E(left), E(y - h / 2), E(w), E(h))
        frame(tb.text_frame)
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if anchor == 'middle' else PP_ALIGN.LEFT
        set_runs(p, s, size, bold, color)
        if name:
            tb.name = name
        return tb

    def box(self, x, y, w, h, fill=None, line=None, lw=2, r=0, dash=False, shadow=False, shape=None, name=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
        sp = self.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
        strip_style(sp)
        if r and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r / min(w, h))
        if fill:
            sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = rgb(line); sp.line.width = Emu(int(lw * PX))
            if dash:
                ln = sp.line._get_or_add_ln()
                etree.SubElement(ln, qn('a:prstDash'), val='dash')
        else:
            sp.line.fill.background()
        if shadow:
            add_shadow(sp)
        frame(sp.text_frame)
        if name:
            sp.name = name
        return sp

    def pill(self, s, cx, cy, size=18, fill=C['navy'], color=C['white'], line=None, padX=20, h=38, bold=True, shadow=False):
        w = measure(s, size, bold) + padX * 2
        sp = self.box(cx - w / 2, cy - h / 2, w, h, fill=fill, line=line, r=h / 2, shadow=shadow, name=f'标签 {s}')
        p = sp.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_runs(p, s, size, bold, color)
        return sp

    def pill_row(self, labels, cx, cy, gap=8, **kw):
        size, bold, padX = kw.get('size', 18), kw.get('bold', True), kw.get('padX', 20)
        ws = [measure(s, size, bold) + padX * 2 for s in labels]
        x = cx - (sum(ws) + gap * (len(labels) - 1)) / 2
        for s, w in zip(labels, ws):
            self.pill(s, x + w / 2, cy, **kw)
            x += w + gap

    def arrow(self, x1, y1, x2, y2, color=C['blue'], lw=3.5, dash=False, ctrl=None):
        if ctrl is None:
            sp = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
            strip_style(sp)
        else:
            sp = self._quad(x1, y1, ctrl, x2, y2)
        sp.line.color.rgb = rgb(color)
        sp.line.width = Emu(int(lw * PX))
        ln = sp.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, qn('a:prstDash'), val='dash')
        etree.SubElement(ln, qn('a:round'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
        sp.name = '箭头'
        return sp

    def _quad(self, x1, y1, c, x2, y2):
        xs, ys = (x1, c[0], x2), (y1, c[1], y2)
        L, T = min(xs), min(ys)
        W, H = max(max(xs) - L, 1), max(max(ys) - T, 1)
        sp = self.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(L), E(T), E(W), E(H))
        strip_style(sp)
        sp.fill.background()
        spPr = sp._element.spPr
        spPr.remove(spPr.find(qn('a:prstGeom')))
        pt = lambda x, y: f'<a:pt x="{int((x - L) * PX)}" y="{int((y - T) * PX)}"/>'
        geom = etree.fromstring(
            f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
            f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst><a:path w="{int(W * PX)}" h="{int(H * PX)}" fill="none">'
            f'<a:moveTo>{pt(x1, y1)}</a:moveTo><a:quadBezTo>{pt(*c)}{pt(x2, y2)}</a:quadBezTo></a:path></a:pathLst></a:custGeom>')
        spPr.insert(1, geom)
        return sp

    def line(self, x1, y1, x2, y2, color=C['line'], lw=2, dash=False):
        sp = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        strip_style(sp)
        sp.line.color.rgb = rgb(color); sp.line.width = Emu(int(lw * PX))
        if dash:
            etree.SubElement(sp.line._get_or_add_ln(), qn('a:prstDash'), val='dash')
        return sp

    def icon(self, key, x, y, w, h, name):
        return self.svg.place(self.shapes, key, x, y, w, h, name)
ST = 'stroke-linejoin="round" stroke-linecap="round"'


def svg_robot(accent):
    n = '#' + C['navy']
    a = '#' + accent
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 120">'
            f'<g stroke="{n}" stroke-width="3.2" {ST}>'
            f'<line x1="50" y1="10" x2="50" y2="24"/><circle cx="50" cy="8" r="6" fill="{a}"/>'
            f'<rect x="6" y="38" width="10" height="24" rx="4" fill="{a}"/><rect x="84" y="38" width="10" height="24" rx="4" fill="{a}"/>'
            f'<rect x="13" y="24" width="74" height="56" rx="19" fill="#FFFFFF"/>'
            f'<path d="M18,118 Q18,86 50,86 Q82,86 82,118 Z" fill="#FFFFFF"/></g>'
            f'<rect x="24" y="35" width="52" height="32" rx="12" fill="{n}"/>'
            f'<circle cx="38" cy="51" r="5.5" fill="#9FD4FF"/><circle cx="62" cy="51" r="5.5" fill="#9FD4FF"/></svg>')


def svg_doc(stroke, line_color, dashed=False):
    s, lc = '#' + stroke, '#' + line_color
    d = ' stroke-dasharray="7 6"' if dashed else ''
    dl = ' stroke-dasharray="6 7"' if dashed else ''
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 80 100"><g fill="none" stroke-width="3.2" {ST}>'
            f'<path d="M4,4 H58 L76,22 V96 H4 Z" fill="#FFFFFF" stroke="{s}"{d}/><path d="M58,4 V22 H76" stroke="{s}"/>'
            f'<line x1="16" y1="44" x2="64" y2="44" stroke="{lc}" stroke-width="4.5"{dl}/>'
            f'<line x1="16" y1="62" x2="64" y2="62" stroke="{lc}" stroke-width="4.5"{dl}/>'
            f'<line x1="16" y1="80" x2="42" y2="80" stroke="{lc}" stroke-width="4.5"{dl}/></g></svg>')


def svg_lock(color):
    c = '#' + color
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 30 34"><g stroke="{c}" stroke-width="3" {ST}>'
            f'<path d="M7,15 V10 A8,8 0 0 1 23,10 V15" fill="none"/><rect x="3" y="14" width="24" height="18" rx="4" fill="#FFFFFF"/></g>'
            f'<circle cx="15" cy="23" r="2.6" fill="{c}"/></svg>')


def svg_eye(color):
    c = '#' + color
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 36 22"><path d="M2,11 Q18,-3 34,11 Q18,25 2,11 Z" fill="#FFFFFF" '
            f'stroke="{c}" stroke-width="3" {ST}/><circle cx="18" cy="11" r="5" fill="{c}"/></svg>')


def svg_check(color):
    c = '#' + color
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 34 34"><circle cx="17" cy="17" r="15" fill="{c}" stroke="#FFFFFF" stroke-width="3"/>'
            f'<path d="M10,17 L15,22 L24,12" fill="none" stroke="#FFFFFF" stroke-width="3.6" {ST}/></svg>')


ICONS = {
    'robot_blue': svg_robot(C['blue']), 'robot_grey': svg_robot(C['neutral']),
    'doc_grey': svg_doc(C['neutral'], C['line']), 'doc_seed': svg_doc(C['rose'], C['rose'], dashed=True),
    'lock_navy': svg_lock(C['navy']), 'lock_grey': svg_lock(C['neutral']),
    'eye_blue': svg_eye(C['blue']), 'check_blue': svg_check(C['blue']),
}


class SvgStore:
    """Embeds each icon once as image/svg+xml plus a PNG fallback rendered by Chrome."""

    def __init__(self, prs, tmp):
        self.prs, self.tmp, self.cache, self.parts = prs, Path(tmp), {}, {}

    def _png(self, key):
        svg = ICONS[key]
        vb = [float(v) for v in re.search(r'viewBox="([^"]+)"', svg).group(1).split()]
        scale = 600 / max(vb[2], vb[3])
        w, h = int(vb[2] * scale), int(vb[3] * scale)
        html = self.tmp / f'{key}.html'
        html.write_text(f'<html><body style="margin:0;background:transparent">'
                        f'<img src="data:image/svg+xml;charset=utf-8,{quote(svg)}" width="{w}" height="{h}"></body></html>', encoding='utf-8')
        out = self.tmp / f'{key}.png'
        subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--default-background-color=00000000',
                        f'--window-size={w},{h}', f'--screenshot={out}', html.as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
        return out

    def place(self, shapes, key, x, y, w, h, name):
        pic = shapes.add_picture(str(self._png_cached(key)), E(x), E(y), E(w), E(h))
        pic.name = name
        if key not in self.parts:
            self.parts[key] = Part(PackURI(f'/ppt/media/icon-{key}.svg'), 'image/svg+xml', self.prs.part.package, ICONS[key].encode('utf-8'))
        rId = shapes.part.relate_to(self.parts[key], RT.IMAGE)
        blip = pic._element.find('.//' + qn('a:blip'))
        ext_lst = etree.SubElement(blip, qn('a:extLst'))
        ext = etree.SubElement(ext_lst, qn('a:ext'), uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        svgb = etree.SubElement(ext, '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
                                nsmap={'asvg': 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
        svgb.set(qn('r:embed'), rId)
        return pic

    def _png_cached(self, key):
        if key not in self.cache:
            self.cache[key] = self._png(key)
        return self.cache[key]
OX, OY = 84, 136  # illustration origin on the slide


def build_left(cv):
    g = cv.group('现有 Agent 评测')
    X = lambda v: OX + v
    Y = lambda v: OY + v
    g.box(X(0), Y(14), 430, 506, fill=C['zone'], line=C['line'], r=18, name='左侧分区')
    g.pill('现有 Agent 评测', X(215), Y(14), size=21, fill=C['neutral'], h=44, padX=26, shadow=True)
    g.pill_row(['SWE-bench', 'GAIA', 'WebArena', 'τ-bench'], X(215), Y(78), size=15, fill=C['white'], color=C['grey'], line=C['line'], padX=12, h=30)
    g.box(X(28), Y(122), 374, 258, fill=C['white'], line=C['neutral'], lw=2.5, r=16, dash=True, name='固定 Harness 框')
    g.pill('固定 Harness', X(215), Y(122), size=16, fill=C['white'], color=C['grey'], line=C['neutral'], h=32, padX=14)
    g.icon('lock_grey', X(380), Y(110), 21, 24, '锁')
    g.pill_row(['tools', 'context', 'memory', 'loop', 'verify'], X(215), Y(168), size=14, fill=C['chipGrey'], color=C['grey'], padX=10, h=28, bold=False)
    g.box(X(54), Y(208), 128, 128, fill=C['pale'], shape=MSO_SHAPE.OVAL, shadow=True, name='底托')
    g.icon('robot_grey', X(72), Y(216), 92, 110, 'Agent 图标')
    g.text('Agent', X(118), Y(356), size=16, bold=True, color=C['grey'])
    g.arrow(X(196), Y(272), X(252), Y(272), color=C['neutral'])
    g.icon('doc_grey', X(266), Y(228), 74, 92, '任务文档')
    g.text('单个任务', X(303), Y(342), size=16, color=C['grey'])
    g.arrow(X(303), Y(384), X(303), Y(418), color=C['neutral'])
    g.pill('被评对象：答案', X(215), Y(452), size=20, fill=C['white'], color=C['navy'], line=C['line'], h=44, padX=22, shadow=True)
    g.text('外壳只是固定的实验配置', X(215), Y(496), size=16, color=C['grey'])


def build_creation(cv):
    g = cv.group('RQ1 Creation')
    X = lambda v: OX + v
    Y = lambda v: OY + v
    g.pill('RQ1', X(566), Y(70), size=15, fill=C['blue'], h=28, padX=12)
    g.text('Creation · 从弱种子构建完整外壳', X(596), Y(70), size=19, bold=True, color=C['navy'], anchor='start')
    g.box(X(560), Y(118), 116, 116, fill=C['paleRose'], shape=MSO_SHAPE.OVAL, shadow=True, name='底托')
    g.icon('doc_seed', X(588), Y(138), 60, 76, '弱种子图标')
    g.text('弱种子 H_{seed}', X(618), Y(252), size=17, bold=True, color=C['rose'])
    g.text('无循环 · 无规划 · 无验证', X(618), Y(276), size=14, color=C['grey'])
    g.arrow(X(684), Y(176), X(734), Y(176))
    g.box(X(742), Y(118), 116, 116, fill=C['pale'], shape=MSO_SHAPE.OVAL, shadow=True, name='底托')
    g.icon('robot_blue', X(760), Y(124), 80, 96, '创建者图标')
    g.text('创建者 L_C', X(800), Y(252), size=17, bold=True, color=C['navy'])
    g.text('读 · 改 · 自测', X(800), Y(276), size=14, color=C['grey'])
    g.text('开发', X(891), Y(154), size=15, bold=True, color=C['blue'])
    g.arrow(X(866), Y(176), X(916), Y(176))
    g.box(X(928), Y(110), 214, 134, fill=C['white'], line=C['line'], r=16, shadow=True, name='H0 模块卡片')
    for i, (k, zh, col) in enumerate(MODS):
        x, y = X(944 + (i % 3) * 64), Y(124 + (i // 3) * 56)
        tile = g.box(x, y, 54, 48, fill=col, r=10, name=f'模块 {k}')
        tf = tile.text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 0.85
        set_runs(p, k, 21, True, C['white'])
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.line_spacing = 0.85
        set_runs(p2, zh, 10, True, C['white'])
    g.icon('check_blue', X(1121), Y(101), 30, 30, '对勾')
    g.text('可运行 H0', X(1035), Y(268), size=17, bold=True, color=C['navy'])
    g.box(X(1178), Y(112), 230, 130, fill=C['paleGold'], line=C['goldLine'], r=14, name='可见性卡片')
    g.icon('eye_blue', X(1192), Y(137), 36, 22, '可见')
    g.text('1–3 个开发样例可见', X(1236), Y(148), size=16, anchor='start')
    g.line(X(1196), Y(177), X(1390), Y(177), color=C['goldLine'])
    g.icon('lock_navy', X(1197), Y(191), 26, 30, '锁')
    g.text('隐藏评测集保留', X(1236), Y(206), size=16, bold=True, color=C['navy'], anchor='start')
def build_evolution(cv):
    g = cv.group('RQ2 Evolution')
    X = lambda v: OX + v
    Y = lambda v: OY + v
    g.pill('RQ2', X(566), Y(336), size=15, fill=C['blue'], h=28, padX=12)
    g.text('Evolution · 用执行反馈持续迭代', X(596), Y(336), size=19, bold=True, color=C['navy'], anchor='start')
    vers = [('H0', '来自 Creation', ['loop', 'tools']), ('H1', '+ 上下文压缩', ['retry ×3']), ('Hn', '+ 记忆 / 验证门', ['memory', 'verify'])]
    for i, (h, d, tags) in enumerate(vers):
        x, y, dark = X(548 + i * 214), Y(372), i == 2
        g.box(x, y, 168, 100, fill=C['navy'] if dark else C['white'], line=C['navy'] if dark else C['line'], r=14, shadow=True, name=f'版本卡片 {h}')
        g.text(h, x + 34, y + 30, size=26, bold=True, color=C['white'] if dark else C['navy'])
        g.text(d, x + 64, y + 32, size=14, color=C['pale'] if dark else C['grey'], anchor='start')
        g.pill_row(tags, x + 84, y + 70, size=13, fill='2B5A8C' if dark else C['pale'], color=C['white'] if dark else C['navy'], padX=10, h=26)
        if i < 2:
            g.arrow(x + 176, y + 50, x + 206, y + 50)
    g.pill('隐藏回退？', X(1164), Y(364), size=13, fill=C['rose'], h=24, padX=10)
    g.arrow(X(1196), Y(422), X(1236), Y(422))
    g.box(X(1248), Y(362), 160, 120, fill=C['white'], line=C['line'], r=14, shadow=True, name='反馈集卡片')
    g.text('反馈集', X(1328), Y(384), size=15, bold=True, color=C['navy'])
    for i, (lab, v, col) in enumerate([('SWE-Pro 100', 0.58, C['blue']), ('Terminal 89', 0.74, C['gold'])]):
        y = Y(410 + i * 34)
        g.text(lab, X(1262), y, size=12, color=C['grey'], anchor='start')
        g.box(X(1262), y + 9, 132, 8, fill=C['pale'], r=4, name='进度底')
        g.box(X(1262), y + 9, 132 * v, 8, fill=col, r=4, name='进度')
    g.arrow(X(1328), Y(488), X(846), Y(480), color=C['gold'], lw=3, dash=True, ctrl=(X(1090), Y(548)))
    g.text('读反馈 → 诊断 → 修改 → 冻结新版本', X(1088), Y(494), size=14, bold=True, color=C['goldText'])
    g.icon('check_blue', X(555), Y(487), 18, 18, '对勾')
    g.text('每个版本冻结、可运行', X(580), Y(496), size=14, color=C['grey'], anchor='start')


def build_slide(prs, store):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cv = Canvas(slide.shapes, store)
    cv.text('HarnessDev 把评测单元从“任务答案”换成“可运行的外壳”', 84, 75, size=40, bold=True, color=C['navy'], anchor='start', name='标题')
    build_left(cv)
    arrow = cv.box(OX + 446, OY + 238, 58, 56, fill=C['gold'], shape=MSO_SHAPE.RIGHT_ARROW, name='过渡箭头')
    arrow.adjustments[0], arrow.adjustments[1] = 0.5, 0.45
    rg = cv.group('HarnessDev')
    rg.box(OX + 518, OY + 14, 914, 506, fill=C['zone'], line=C['line'], r=18, name='右侧分区')
    rg.pill('HarnessDev：模型不仅解题，还要构建系统', OX + 975, OY + 14, size=21, fill=C['navy'], h=44, padX=28, shadow=True)
    rg.line(OX + 546, OY + 306, OX + 1404, OY + 306, dash=True)
    build_creation(rg)
    build_evolution(rg)
    # formula band
    cv.box(84, 676, 1432, 116, fill=C['pale'], name='公式条')
    cv.text('(L_C, D)  →  H          (H, L_E, x)  →  y  →  score (J)', 120, 718, size=30, bold=True, color=C['navy'], anchor='start', name='公式')
    cv.text('创建者 L_C 在开发环境 D 中产出外壳 H；H 冻结后，再交给执行者 L_E 运行下游任务 x', 120, 766, size=21, anchor='start', name='公式说明')
    cv.text('3', 1490, 840, size=18, color=C['grey'], anchor='start', name='页码')
    return slide
def build_library(prs, store):
    """A component sheet: copy-paste these into any slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cv = Canvas(slide.shapes, store)
    cv.text('组件库：复制粘贴到任意页面使用', 84, 75, size=40, bold=True, color=C['navy'], anchor='start', name='标题')
    cv.text('文字、标签、卡片、箭头为原生形状，可直接改字改色；图标为 SVG，可在“图形格式”里改色或右键“转换为形状”', 84, 130, size=18, color=C['grey'], anchor='start')
    cv.text('图标（SVG）', 84, 190, size=22, bold=True, color=C['navy'], anchor='start')
    icons = [('robot_blue', 80, 96, '机器人·蓝'), ('robot_grey', 80, 96, '机器人·灰'), ('doc_seed', 64, 80, '种子文档'), ('doc_grey', 64, 80, '文档·灰'),
             ('lock_navy', 34, 40, '锁·深蓝'), ('lock_grey', 34, 40, '锁·灰'), ('eye_blue', 48, 30, '可见'), ('check_blue', 40, 40, '对勾')]
    for i, (k, w, h, label) in enumerate(icons):
        cx = 150 + i * 170
        cv.box(cx - 62, 220, 124, 124, fill=C['pale'], shape=MSO_SHAPE.OVAL, shadow=True, name='底托')
        cv.icon(k, cx - w / 2, 282 - h / 2, w, h, label)
        cv.text(label, cx, 366, size=14, color=C['grey'])
    cv.text('标签与卡片（原生形状）', 84, 420, size=22, bold=True, color=C['navy'], anchor='start')
    cv.pill('分区标题', 170, 476, size=21, fill=C['navy'], h=44, padX=26, shadow=True)
    cv.pill('RQ1', 330, 476, size=15, fill=C['blue'], h=28, padX=12)
    cv.pill('隐藏回退？', 440, 476, size=13, fill=C['rose'], h=24, padX=10)
    cv.pill('tools', 560, 476, size=14, fill=C['chipGrey'], color=C['grey'], padX=10, h=28, bold=False)
    cv.pill('memory', 660, 476, size=13, fill=C['pale'], color=C['navy'], padX=10, h=26)
    cv.pill('被评对象：答案', 840, 476, size=20, fill=C['white'], color=C['navy'], line=C['line'], h=44, padX=22, shadow=True)
    for i, (k, zh, col) in enumerate(MODS):
        tile = cv.box(1040 + i * 64, 452, 54, 48, fill=col, r=10, name=f'模块 {k}')
        p = tile.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 0.85
        set_runs(p, k, 21, True, C['white'])
        p2 = tile.text_frame.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.line_spacing = 0.85
        set_runs(p2, zh, 10, True, C['white'])
    card = cv.box(84, 540, 230, 110, fill=C['white'], line=C['line'], r=14, shadow=True, name='白卡片')
    cv.box(344, 540, 230, 110, fill=C['navy'], r=14, shadow=True, name='深色卡片')
    cv.box(604, 540, 230, 110, fill=C['paleGold'], line=C['goldLine'], r=14, name='金色提示卡')
    cv.box(864, 540, 230, 110, fill=C['zone'], line=C['line'], r=18, name='分区底板')
    cv.box(1124, 540, 230, 110, fill=C['white'], line=C['neutral'], lw=2.5, r=16, dash=True, name='虚线框')
    cv.text('箭头（原生线条）', 84, 704, size=22, bold=True, color=C['navy'], anchor='start')
    cv.arrow(100, 760, 220, 760)
    cv.arrow(260, 760, 380, 760, color=C['neutral'])
    cv.arrow(420, 770, 640, 770, color=C['gold'], lw=3, dash=True, ctrl=(530, 720))
    a = cv.box(690, 734, 58, 56, fill=C['gold'], shape=MSO_SHAPE.RIGHT_ARROW, name='过渡箭头')
    a.adjustments[0], a.adjustments[1] = 0.5, 0.45


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else 'slide3-native.pptx')
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    with tempfile.TemporaryDirectory() as tmp:
        store = SvgStore(prs, tmp)
        build_slide(prs, store)
        build_library(prs, store)
        prs.save(out)
    print(f'saved {out}')


if __name__ == '__main__':
    main()
